import csv
import io
from decimal import Decimal

import qrcode
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from qrcode.constants import ERROR_CORRECT_H
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from app.models.entities import AlertaPedido, Producto
from app.schemas.inventory import ProductoCreate

MAX_IMPORT_BYTES = 10 * 1024 * 1024
CATALOG_HEADERS = ["codigo_qr", "nombre", "categoria", "ubicacion", "unidad_medida", "stock_minimo", "stock_maximo"]
OPTIONAL_CATALOG_HEADERS = ["descripcion", "foto_url", "stock_inicial"]


def excel_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def read_catalog(content: bytes, filename: str) -> tuple[list[tuple[int, ProductoCreate, Decimal]], list[dict]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("El archivo supera el límite de 10 MB")
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
        first_row = 2
    elif filename.lower().endswith((".xlsx", ".xlsm")):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        iterator = sheet.iter_rows(values_only=True)
        headings = [str(v).strip() if v is not None else "" for v in next(iterator, ())]
        rows = [dict(zip(headings, values)) for values in iterator]
        first_row = 2
    else:
        raise ValueError("Formato no soportado; use .csv o .xlsx")
    if not rows and not (filename.lower().endswith(".csv") and content):
        return [], []
    parsed, errors = [], []
    for line, row in enumerate(rows, first_row):
        try:
            normalized = {str(k).strip(): v for k, v in row.items() if k is not None}
            missing = [header for header in CATALOG_HEADERS if header not in normalized]
            if missing:
                raise ValueError("Faltan columnas: " + ", ".join(missing))
            product_values = {key: normalized[key] for key in CATALOG_HEADERS}
            product_values.update({key: normalized.get(key, "") or "" for key in OPTIONAL_CATALOG_HEADERS[:2]})
            item = ProductoCreate.model_validate(product_values)
            initial = Decimal(str(normalized.get("stock_inicial", 0) or 0))
            if not initial.is_finite() or initial < 0 or initial != initial.to_integral_value():
                raise ValueError("stock_inicial debe ser un entero no negativo")
            parsed.append((line, item, initial))
        except Exception as exc:
            errors.append({"fila": line, "error": str(exc)[:500]})
    return parsed, errors


def demo_catalog_xlsx() -> bytes:
    from app.demo_catalog import DEMO_PRODUCTS

    wb = Workbook(); ws = wb.active; ws.title = "Catalogo de ejemplo"
    headers = CATALOG_HEADERS + OPTIONAL_CATALOG_HEADERS
    ws.append(headers)
    for product in DEMO_PRODUCTS:
        ws.append([product.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO(); wb.save(output)
    return output.getvalue()


def suggestions_xlsx(alerts: list[AlertaPedido], products: dict[int, Producto]) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Pedidos sugeridos"
    ws.append(["Código QR", "Producto", "Categoría", "Ubicación", "Unidad", "Stock actual", "Mínimo", "Máximo", "Cantidad sugerida", "Estado"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="14532D")
    for alert in alerts:
        p = products.get(alert.producto_id)
        if p:
            ws.append([excel_safe(p.codigo_qr), excel_safe(p.nombre), excel_safe(p.categoria), excel_safe(p.ubicacion), excel_safe(p.unidad_medida),
                       float(p.stock_actual), float(p.stock_minimo), float(p.stock_maximo),
                       float(alert.cantidad_sugerida), alert.estado.value])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(34, max(12, max(len(str(c.value or "")) for c in col) + 2))
    out = io.BytesIO(); wb.save(out); return out.getvalue()


def suggestions_pdf(alerts: list[AlertaPedido], products: dict[int, Producto]) -> bytes:
    out = io.BytesIO()
    data = [["Código", "Producto", "Stock", "Mínimo", "Máximo", "Sugerido"]]
    for a in alerts:
        p = products.get(a.producto_id)
        if p:
            data.append([p.codigo_qr, p.nombre[:32], str(p.stock_actual), str(p.stock_minimo), str(p.stock_maximo), str(a.cantidad_sugerida)])
    table = Table(data, colWidths=[72, 180, 55, 55, 55, 72], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#14532d")),
                               ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .4, colors.grey),
                               ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 8),
                               ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f0fdf4")])]))
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(out, pagesize=letter, title="Pedidos sugeridos de bodega",
                            leftMargin=30, rightMargin=30, topMargin=34, bottomMargin=34)
    doc.build([Paragraph("Pedidos sugeridos", styles["Title"]), Spacer(1, 12), table])
    return out.getvalue()


def qr_labels_pdf(products: list[Producto]) -> bytes:
    out = io.BytesIO(); doc = canvas.Canvas(out, pagesize=letter)
    page_w, page_h = letter
    columns, rows = 3, 8
    cell_w, cell_h = page_w / columns, page_h / rows
    for idx, product in enumerate(products):
        if idx and idx % (columns * rows) == 0:
            doc.showPage()
        cell = idx % (columns * rows); col, row = cell % columns, cell // columns
        x, y = col * cell_w, page_h - (row + 1) * cell_h
        qr = qrcode.QRCode(error_correction=ERROR_CORRECT_H, box_size=6, border=2)
        qr.add_data(product.codigo_qr); qr.make(fit=True)
        image = qr.make_image(fill_color="black", back_color="white")
        qr_side = min(cell_h - 34, cell_w - 16)
        doc.drawImage(ImageReader(image.get_image()), x + (cell_w - qr_side) / 2, y + 24, qr_side, qr_side)
        doc.setFont("Helvetica-Bold", 8); doc.drawCentredString(x + cell_w / 2, y + 12, product.codigo_qr)
        doc.setFont("Helvetica", 7); doc.drawCentredString(x + cell_w / 2, y + 3, product.nombre[:34])
        doc.setStrokeColor(colors.lightgrey); doc.rect(x + 3, y + 3, cell_w - 6, cell_h - 6)
    doc.save(); return out.getvalue()
